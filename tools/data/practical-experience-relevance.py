"""
Evaluation-Matrix criterion #3 — «المجال» relevance.

Reads the authoritative classification workbook
(`docs/inputs/practical-experience-relevance.xlsx`) and generates:

  backend/…/Persistence/PracticalExperienceRelevance.cs
      the criterion's option→percentage-point table.
  frontend/…/shared/content/domainCatalogueStatus.ts
      the codes that are NOT valid domains, so a NEW application's dropdown can
      stop offering them without touching the historical master list.

Criterion #3 scores «المجال» — Section 1's own field (Form 1 #16 →
`cmpt.JobFamily`, the source column the matrix names). Notion renamed the
criterion to «المجال» on 2026-09-22 and the owner confirmed it: «في الخبرة
العملية لا يوجد مجال للخبرة المجال فقط في APPLICATION FORM». The classification
is COMMON — one answer serves Trainer, Consultant, Content Developer and
Question Writer alike.

Three controlled classifications, and the third is the point:

  RELATED                     → 0.1  (10.00 points)
  NOT_RELATED                 → 0    (0.00 points)
  MASTER_DATA_REVIEW_REQUIRED → OMITTED from the table

A value that is not a professional domain at all — a tool, a certification, a
methodology, a programme, a service type, a catch-all — is not «not related».
Recording it as NOT_RELATED would state a business decision nobody made. It is
left out of the table instead, and the look-up is `strict`, so an applicant who
picked one is reported UNRESOLVED and the problem stays diagnosable.

⚠️ NOTHING here is inferred. No string matching, no fuzzy matching, no AI. A
value pays only what the workbook's own classification column says.

Run from the repository root:
    python tools/data/practical-experience-relevance.py extract
"""

from __future__ import annotations

import io
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "docs" / "inputs" / "practical-experience-relevance.xlsx"
DOMAINS_TS = (
    ROOT / "frontend" / "src" / "shared" / "content"
    / "specializationDomains.ts"
)
OUT_CS = (
    ROOT / "backend" / "src" / "ExpertHub.Infrastructure"
    / "Persistence" / "PracticalExperienceRelevance.cs"
)
OUT_TS = (
    ROOT / "frontend" / "src" / "shared" / "content"
    / "domainCatalogueStatus.ts"
)

SHEET = "المجال"
RELATED = "RELATED"
NOT_RELATED = "NOT_RELATED"
INVALID = "MASTER_DATA_REVIEW_REQUIRED"

RELATED_POINTS = "0.1"
NOT_RELATED_POINTS = "0"


def read_domains() -> dict[str, str]:
    """The approved list, parsed from the module both forms already import."""
    source = io.open(DOMAINS_TS, encoding="utf8").read()
    rows = re.findall(r"\{ value: '([^']+)', labelAr: '((?:[^'\\]|\\.)*)'", source)
    if len(rows) != 147:
        raise SystemExit(
            f"expected the approved 147-value list, found {len(rows)} in "
            f"{DOMAINS_TS.relative_to(ROOT)}"
        )
    return {code: label.replace("\\'", "'") for code, label in rows}


def chunk(pairs: list[str], per_line: int = 6) -> str:
    lines = [
        ",".join(pairs[i : i + per_line]) + ("," if i + per_line < len(pairs) else "")
        for i in range(0, len(pairs), per_line)
    ]
    return "\n        + ".join(f'"{line}"' for line in lines) or '""'


def extract() -> None:
    if not WORKBOOK.exists():
        raise SystemExit(f"{WORKBOOK.relative_to(ROOT)} does not exist")
    domains = read_domains()
    sheet = load_workbook(WORKBOOK, data_only=True)[SHEET]

    points: dict[str, str] = {}
    invalid: list[tuple[str, str]] = []
    seen: set[str] = set()
    related = 0

    for row in sheet.iter_rows(min_row=2, values_only=True):
        code = str(row[1]).strip() if row[1] is not None else ""
        if not code.startswith("dom-"):
            continue
        if code not in domains:
            raise SystemExit(
                f"{code!r} is not a value of the approved list — the workbook and "
                f"{DOMAINS_TS.name} have diverged"
            )
        if code in seen:
            raise SystemExit(f"{code!r} appears more than once in the workbook")
        seen.add(code)

        classification = str(row[4]).strip() if row[4] is not None else ""
        if classification not in (RELATED, NOT_RELATED, INVALID):
            raise SystemExit(
                f"{code}: unrecognised classification {classification!r} — expected "
                f"{RELATED}, {NOT_RELATED} or {INVALID}"
            )
        if classification == INVALID:
            # Omitted from the table ON PURPOSE. See the module docstring.
            invalid.append((code, str(row[7]).strip() if row[7] is not None else ""))
            continue
        points[code] = RELATED_POINTS if classification == RELATED else NOT_RELATED_POINTS
        related += classification == RELATED

    missing = set(domains) - seen
    if missing:
        raise SystemExit(
            f"{len(missing)} approved values are absent from the workbook "
            f"(e.g. {sorted(missing)[:3]})"
        )

    classified = len(points)
    not_related = classified - related
    table = chunk([f'\\"{code}\\":{value}' for code, value in points.items()])
    io.open(OUT_CS, "w", encoding="utf8", newline="\n").write(
        f'''namespace ExpertHub.Infrastructure.Persistence;

/// <summary>
/// Criterion #3 «المجال» — the option→percentage-point table.
/// </summary>
/// <remarks>
/// ⚠️ GENERATED — do not edit by hand. Produced by
/// `tools/data/practical-experience-relevance.py extract` from
/// `docs/inputs/practical-experience-relevance.xlsx`, the business's own
/// classification. The keys are the approved 147-value «مجال التخصص» list the
/// application form's «المجال» field already offers, so the value a person
/// picks and the table it is scored against come from one master list.
///
/// The matrix's table: «مجال ليس له صلة 0 · مجال ذو صلة 0.1», max 10%. The
/// classification is COMMON — one answer serves all four services (approved
/// policy, 2026-09-28).
///
/// <b>{related} RELATED · {not_related} NOT_RELATED · {len(invalid)} omitted.</b>
///
/// The {len(invalid)} omitted values are not professional domains at all —
/// tools, certifications, methodologies, programmes, service types, a
/// catch-all. They are NOT recorded as «not related», because that would state
/// a relevance decision nobody made. They are left OUT of the table, and the
/// rule is `strict`, so an applicant who picked one is reported UNRESOLVED and
/// the master-data problem stays diagnosable. All 147 codes are preserved in
/// the form's own list for historical compatibility.
/// </remarks>
internal static class PracticalExperienceRelevance
{{
    /// <summary>Every approved value carries a business classification.</summary>
    internal const bool IsComplete = true;

    /// <summary>Values the business has ruled on, and that can therefore score.</summary>
    internal const int ClassifiedCount = {classified};

    /// <summary>Values that are not domains, and so resolve to nothing.</summary>
    internal const int MasterDataReviewCount = {len(invalid)};

    internal const string ScoreRule =
        "{{\\"kind\\":\\"option\\",\\"strict\\":true,\\"points\\":{{"
        + {table}
        + "}}}}";
}}
'''
    )

    entries = "\n".join(
        f"  // {domains[code]} — {problem or 'master-data review'}\n  '{code}',"
        for code, problem in invalid
    )
    io.open(OUT_TS, "w", encoding="utf8", newline="\n").write(
        f'''/**
 * Domain codes that are NOT valid professional domains — tools,
 * certifications, methodologies, programmes, service types and a catch-all
 * that reached the legacy «مجال التخصص» master list.
 *
 * ⚠️ GENERATED — do not edit by hand. Produced by
 * `tools/data/practical-experience-relevance.py extract` from the
 * business's own classification workbook.
 *
 * They are **kept in the master list** ({len(invalid)} of 147): a historical
 * application that selected one must stay readable forever, and no code is
 * renamed, merged or deleted. This list exists only so a NEW application's
 * «المجال» dropdown can stop offering them — the form schema is versioned and
 * each version carries its own frozen copy of the options, so filtering the
 * current version leaves every earlier one untouched.
 *
 * Evaluation-Matrix criterion #3 reports these as UNRESOLVED rather than
 * scoring them, which is the same statement made on the scoring side.
 *
 * Cleanup backlog: `EXPERT-HUB-DOMAIN-MASTER-CLEANUP`.
 */

export const INACTIVE_DOMAIN_CODES: readonly string[] = [
{entries}
];

/** The catalogue a NEW application should offer — the master list minus those. */
export function activeDomains<T extends {{ readonly value: string }}>(
  catalogue: readonly T[]
): readonly T[] {{
  return catalogue.filter((option) => !INACTIVE_DOMAIN_CODES.includes(option.value));
}}
'''
    )

    print(f"wrote {OUT_CS.relative_to(ROOT)}")
    print(f"wrote {OUT_TS.relative_to(ROOT)}")
    print(f"  RELATED {related} / NOT_RELATED {not_related} / omitted {len(invalid)}")
    print(f"  total {classified + len(invalid)}")


if __name__ == "__main__":
    if (sys.argv[1] if len(sys.argv) > 1 else "") == "extract":
        extract()
    else:
        raise SystemExit(__doc__)
